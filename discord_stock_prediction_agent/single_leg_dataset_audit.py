"""Audit single-leg option parser output against JSON/XLSX expectations."""
from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from .options_parser import classify_and_parse


def _json_rows(path: Path) -> list[dict[str, Any]]:
    rows = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(rows, list):
        raise ValueError("Expected a JSON list of test cases.")
    return rows


def _xlsx_rows(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - optional audit dependency
        raise RuntimeError("openpyxl is required to audit an XLSX file.") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Single-Leg Options"]
    rows: list[dict[str, Any]] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row or not row[3]:
            continue
        rows.append(
            {
                "id": row[0],
                "category": row[1],
                "difficulty": row[2],
                "signal": row[3],
                "expected_output": json.loads(str(row[4])),
                "features_tested": row[5],
            }
        )
    return rows


def audit_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    mismatches: list[dict[str, Any]] = []
    for index, row in enumerate(rows, start=1):
        signal = str(row.get("signal") or "")
        routed = classify_and_parse(signal)
        actual = routed.option.semantic_contract if routed.option else None
        expected = row.get("expected_output")
        if actual != expected:
            mismatches.append(
                {
                    "row": index,
                    "id": row.get("id"),
                    "signal": signal,
                    "expected": expected,
                    "actual": actual,
                }
            )
    return mismatches


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("dataset", type=Path, help="Expected-output JSON or XLSX file")
    args = parser.parse_args()
    rows = _xlsx_rows(args.dataset) if args.dataset.suffix.lower() == ".xlsx" else _json_rows(args.dataset)
    mismatches = audit_rows(rows)
    print(
        f"TOTAL {len(rows)} MATCH {len(rows) - len(mismatches)} "
        f"MISMATCH {len(mismatches)}"
    )
    if mismatches:
        print(json.dumps(mismatches, indent=2))
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
